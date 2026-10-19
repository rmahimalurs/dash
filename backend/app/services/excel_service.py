import io
import re
import logging
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

_TEMPLATES = Path(__file__).parent.parent / "templates"


def _tz_label(iso_dt: str) -> str:
    """ET is EDT or EST depending on the date; the offset in the ISO string says
    which. Hardcoding EST labelled every summer window wrong."""
    return "EDT" if "-0400" in iso_dt else "EST"


def _parse_iso(iso_dt: str) -> datetime:
    return datetime.strptime(iso_dt.split(".")[0], "%Y-%m-%dT%H:%M:%S")


def _format_dt(iso_dt: str) -> str:
    """'2026-10-17T09:00:00.000-0400' → '10/17/2026 09:00 AM EDT'"""
    return _parse_iso(iso_dt).strftime(f"%m/%d/%Y %I:%M %p {_tz_label(iso_dt)}")


def _clock(dt: datetime) -> str:
    """Match the template's own style: '09:00 p.m. ET'."""
    return dt.strftime("%I:%M %p ET").replace("AM", "a.m.").replace("PM", "p.m.")


_DURATION_RE = re.compile(r"(\d+)\s*min", re.I)

# Where the task table lives. Anticipated Start/End are columns I and K.
_TASK_ROWS = range(39, 57)


def _schedule_tasks(ws, start_dt: str) -> int:
    """Re-time the task table so it starts when the window does.

    The template ships a hand-built chain of times from a 9 p.m. start and
    nothing updated them, so a 7 a.m. window produced a plan whose header said
    7 a.m. and whose steps said 9 p.m. Each task's Est. Duration is chained
    from the window start instead, and the Anticipated End Time column, which
    the template left empty, is filled in too.
    """
    clock = _parse_iso(start_dt)
    written = 0
    for row in _TASK_ROWS:
        dur = ws[f"F{row}"].value
        if not isinstance(dur, str):
            continue
        m = _DURATION_RE.search(dur)
        if not m:
            continue
        minutes = int(m.group(1))
        _safe_set(ws, f"I{row}", _clock(clock))
        clock += timedelta(minutes=minutes)
        _safe_set(ws, f"K{row}", _clock(clock))
        written += 1
    return written


def _newsletter_date(go_nogo_date: str) -> str:
    """'YYYY-MM-DD' → 'M/D/YYYY' (no leading zeros)"""
    dt = datetime.strptime(go_nogo_date, "%Y-%m-%d")
    return f"{dt.month}/{dt.day}/{dt.year}"


def _wiki_to_plain_lines(description: str) -> list[str]:
    """Strip Jira wiki markup from a description, stopping before runbook links."""
    result = []
    for line in description.split("\n"):
        if "*Link to deployment runbooks:*" in line or "Link to smoke testing plan:" in line:
            break
        if line.startswith(" * "):
            result.append(line[3:].strip().replace("*", ""))
        elif line.startswith(" ** "):
            result.append(line[4:].strip())
        else:
            result.append(line)
    return result


def _to_cms_description(description: str) -> str:
    """Plain text for CMS Change Management form — keeps intro, drops systems impacted + runbooks."""
    lines = _wiki_to_plain_lines(description)
    result = []
    for line in lines:
        if "Systems impacted" in line:
            break
        result.append(line)
    return "\n".join(result).rstrip()


def _to_ir_plan_description(description: str) -> str:
    """Plain text for IR Plan — drops intro, keeps systems impacted, drops runbooks."""
    lines = _wiki_to_plain_lines(description)
    plain = "\n".join(lines)
    if plain.startswith("The below describes"):
        idx = plain.find("\n")
        if idx >= 0:
            plain = plain[idx + 1:].lstrip("\n")
    return plain.rstrip()


from openpyxl.cell.cell import MergedCell


def _safe_set(ws, coord: str, value) -> None:
    """Set a cell value, skipping merged non-anchor cells (read-only in openpyxl)."""
    cell = ws[coord]
    if not isinstance(cell, MergedCell):
        cell.value = value


# Release Manager for the IR plan (row 29). Replaces the template's Patrick Hendren.
_IR_RELEASE_MANAGER = {"name": "Meghana Reddy", "org": "CMS", "email": "meghana.reddy@samtek.io"}

# Keyword → tool for matching an implementation/validation task's text to an owner.
_TASK_TOOL_KEYWORDS = [
    ("jfrog", "jfrog"), ("artifactory", "jfrog"), ("xray", "jfrog"),
    ("cloudbees", "cloudbees"), ("jenkins", "cloudbees"), ("cbj", "cloudbees"),
    ("sonarqube", "sonarqube"), ("sonar", "sonarqube"),
]


def generate_ir_plan(
    mggd_key: str,
    itc_key: str,
    go_nogo_date: str,    # "YYYY-MM-DD"
    start_dt: str,        # ISO e.g. "2026-10-17T09:00:00.000-0400"
    end_dt: str,
    description: str,     # Jira wiki markup
    maintenance_date: str,  # "MM-DD-YYYY" used in filename only
    tool_owners: dict | None = None,  # {"jfrog"/"cloudbees"/"sonarqube": display name}
) -> bytes:
    wb = openpyxl.load_workbook(_TEMPLATES / "ir_plan_template.xlsx")
    ws = wb["Artifact Implementation Plan"]

    ws["K2"] = itc_key or ""
    ws["C3"] = mggd_key
    ws["C5"] = datetime.strptime(go_nogo_date, "%Y-%m-%d")
    ws["C8"] = _format_dt(start_dt)
    ws["K8"] = _format_dt(end_dt)
    ws["C11"] = _to_ir_plan_description(description)

    # Re-time the task table to the real window; see _schedule_tasks.
    scheduled = _schedule_tasks(ws, start_dt)
    if not scheduled:
        logger.warning("IR plan: no task durations found, schedule left as the template had it")

    # Release Manager (row 29) -> Meghana; remove the Dev Pathak POC row (33).
    # The roles table uses merged cells, so write only anchor cells.
    _safe_set(ws, "C29", _IR_RELEASE_MANAGER["name"])
    _safe_set(ws, "E29", _IR_RELEASE_MANAGER["org"])
    _safe_set(ws, "K29", _IR_RELEASE_MANAGER["email"])
    for cell in ws[33]:  # clear the whole Dev Pathak row (name/org/email/hyperlink)
        if not isinstance(cell, MergedCell):
            cell.value = None
            cell.hyperlink = None

    # Implementation/validation task owners (col D, rows 39-53) -> the engineer
    # assigned to that tool's prep ticket, matched by keyword in the task text (col B).
    if tool_owners:
        for row in range(39, 54):
            task = ws[f"B{row}"].value
            if not isinstance(task, str):
                continue
            low = task.lower()
            tool = next((t for kw, t in _TASK_TOOL_KEYWORDS if kw in low), None)
            if tool and tool_owners.get(tool):
                _safe_set(ws, f"D{row}", tool_owners[tool])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_change_management_form(
    start_dt: str,
    end_dt: str,
    description: str,
    go_nogo_date: str,
    maintenance_date: str,  # "MM-DD-YYYY" used in filename only
) -> bytes:
    wb = openpyxl.load_workbook(_TEMPLATES / "change_management_template.xlsx")
    ws = wb["Sheet1"]

    ws["G2"] = _format_dt(start_dt)
    ws["G4"] = _format_dt(end_dt)
    ws["A6"] = _to_cms_description(description)
    ws["A8"] = f"CMS Cloud Weekly Change Newsletter {_newsletter_date(go_nogo_date)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
